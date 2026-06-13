"""Generate v3 Word manual from structured content."""
import datetime, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(root))

doc = Document()

# Title
title = doc.add_heading("Scientra Copilot", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("AI-Powered Research Discovery Platform", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph("Version 1.5 (Storage Layout v3)")
doc.add_paragraph(f"Generated: {datetime.date.today().isoformat()}")
doc.add_paragraph("")

sections = [
    ("1. About", "Scientra Copilot is an AI-powered research literature discovery platform. "
     "It processes PDFs, extracts evidence, builds semantic indexes, and provides an AI chat interface "
     "with traceable citations. All answers are grounded in your literature database."),
    ("2. Quick Start", [
        "Requirements: Python 3.11+, Docker (GROBID), Node.js 18+",
        "Start all: python Scripts/dev_restart.py",
        "Alternative: Scripts/launchers/start_scientra.bat or .ps1",
        "API only: python -m scientra.server (http://127.0.0.1:8710)",
        "Web only: cd web && npm run dev (http://127.0.0.1:3000)",
        "Setup LLM: python Scripts/setup_llm.py",
    ]),
    ("3. Storage Layout v3", [
        "00_Inbox/ - Import staging (recommended: article_bundles)",
        "01_Sources/ - Original papers + supplementary files",
        "02_Parse/ - Parsed intermediates (text, figures, tables)",
        "03_Assets/ - Structured searchable assets",
        "04_Corpus/ - Long-term corpora (writing, reasoning, data, review)",
        "05_Knowledge/ - Knowledge graph + research maps",
        "06_Index/ - Search indexes (LanceDB primary path)",
        "07_Agents/ - Agent workspaces",
        "08_Projects/ - User project workspaces",
        "09_Exports/ - Export outputs",
        "10_System/ - System files (config, logs, migrations, backups, legacy archive)",
    ]),
    ("4. Recommended Import: Article Bundle", [
        "One folder = one article. Place main PDF and supplementary files together.",
        "",
        "Example:",
        "  00_Inbox/article_bundles/new/Article_001/",
        "    main.pdf              - Main paper PDF",
        "    Table_S1.xlsx         - Supplementary table",
        "    Table_S2.csv          - Supplementary data",
        "    Supplementary_Info.pdf - Supplementary material",
        "",
        "Commands:",
        "  python Scripts/process_article_bundles.py --scan",
        "  python Scripts/process_article_bundles.py --process --archive-mode copy",
    ]),
    ("5. folder_explicit Binding", [
        "main.pdf = main paper (auto-detected)",
        "Other files in folder = supplementary files for THAT paper",
        "Match confidence = high (folder-explicit binding)",
        "No filename guessing needed",
        "loose supplementary (files without folder context) require manual review",
        "loose supplementary files are NOT auto-matched as high confidence",
    ]),
    ("6. Chat Query Types", [
        "claim_query - Evidence quality assessment",
        "research_gap_query - Gap identification",
        "method_query - Method analysis",
        "result_query - Result synthesis",
        "supplementary_entity_query - Entity lookup in supplementary data",
        "supplementary_entity_comparison_query - Cross-paper entity comparison",
        "",
        "Three answer modes: Auto (recommended), LLM synthesis, Evidence-only",
    ]),
    ("7. Web Pages", [
        "/chat - AI Chat [OK]",
        "/library - Paper library [OK]",
        "/paper/{id} - Paper detail + Ask this paper [OK]",
        "/evidence - Evidence search [OK]",
        "/research-map - Research topic map [OK]",
        "/hotspots - Trending topics [OK]",
        "/research-gaps - Research gaps [OK]",
        "/knowledge-network - Knowledge network [OK]",
        "/report - Library report [OK]",
        "/import - Import Dashboard [PLANNED]",
        "/assets - Assets Viewer [PLANNED]",
    ]),
    ("8. LanceDB Path", [
        "Primary: 06_Index/vector/lancedb/",
        "Tables:",
        "  evidence_chunks - 1,267 rows",
        "  literature_vectors - 319 rows",
        "  pdf_asset_chunks - 2,648 rows (11 chunk types)",
        "",
        "Old path 04_VectorDB/ has been archived. Do not use.",
    ]),
    ("9. Data Safety", [
        "Files are NEVER auto-deleted. All operations default to copy mode.",
        "Old directories (01_PDF, 03_Evidence, etc.) archived to 10_System/legacy_archive/.",
        "LanceDB backups created before promotion at 10_System/backups/.",
        "All reports use relative paths only. No absolute paths exposed.",
        "LLM API keys stored in Config/llm_config.yaml (gitignored).",
    ]),
    ("10. Current Limitations", [
        "Supplementary PDFs are not OCR-processed.",
        "ZIP files are not auto-extracted.",
        "loose supplementary files require manual confirmation.",
        "Cross-paper value comparisons are indicative, not statistically validated.",
        "Import Dashboard and Assets Viewer are under development.",
    ]),
    ("11. FAQ", [
        "Q: Why Article Bundle instead of single PDF import?",
        "A: Folder-explicit binding ensures supplementary files are correctly linked to their paper.",
        "",
        "Q: What is loose supplementary?",
        "A: Supplementary files without a parent article bundle. They need manual review.",
        "",
        "Q: Where is LanceDB?",
        "A: Primary path is 06_Index/vector/lancedb/. Old 04_VectorDB/ is archived.",
        "",
        "Q: How to delete legacy archive?",
        "A: After confirming all systems work, manually delete 10_System/legacy_archive/.",
    ]),
]

for title_text, content in sections:
    doc.add_heading(title_text, level=1)
    if isinstance(content, str):
        doc.add_paragraph(content)
    elif isinstance(content, list):
        for item in content:
            if item == "":
                doc.add_paragraph("")
            else:
                doc.add_paragraph(item, style="List Bullet")

doc.add_paragraph("")
doc.add_paragraph("Scientra Copilot - From Literature to Discovery.", style="Intense Quote")

out = root / "docs" / "manual" / "Scientra_Copilot_使用手册_v3.docx"
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(out))
print(f"Saved: {out.relative_to(root)}")
print(f"Size: {out.stat().st_size} bytes")
